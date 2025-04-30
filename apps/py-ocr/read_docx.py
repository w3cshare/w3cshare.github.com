import os
from docx import Document

def read_docx(file_path: str) -> str:
    """
    读取Word文档内容

    Args:
        file_path (str): Word文档路径

    Returns:
        str: 文档内容
    """
    try:
        # 检查文件是否存在
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"文件不存在: {file_path}")

        # 检查文件扩展名
        if not file_path.endswith('.docx'):
            raise ValueError("只支持.docx格式的文件")

        # 打开文档
        doc = Document(file_path)
        full_text = []

        # 读取段落
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text)

        # 读取表格
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells]
                full_text.append(' | '.join(row_text))

        return '\n'.join(full_text)

    except Exception as e:
        print(f"读取文档时出错: {str(e)}")
        return ""

def main():
    # 文档路径
    docx_path = "PaddleXOCR技术评审.docx"

    try:
        # 读取文档内容
        content = read_docx(docx_path)
        if content:
            print("文档内容:")
            print("-" * 50)
            print(content)
            print("-" * 50)

            return content
        else:
            print("未能成功读取文档内容")
    except Exception as e:
        print(f"程序执行出错: {str(e)}")

if __name__ == "__main__":
    main()
